import io
import json

import pymupdf
from docx import Document
from PIL import Image, ImageDraw
from pptx import Presentation

from vault_retrieval.extract import Extractor


def test_pdf_text_scans_cache_and_version(tmp_path):
    extractor = Extractor(tmp_path)
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((40, 60), "A source page about Gaussian elimination.")
    image = Image.new("RGB", (1000, 300), "white")
    ImageDraw.Draw(image).text((40, 80), "SCANNED LECTURE EVIDENCE", fill="black", font_size=40)
    image_data = io.BytesIO()
    image.save(image_data, format="PNG")
    page = doc.new_page()
    page.insert_image(page.rect, stream=image_data.getvalue())
    data = doc.tobytes()
    result, hit, key = extractor.extract(data, ".pdf")
    assert not hit
    assert result["sections"][0]["locator"]["page"] == 1
    assert "visual_review_required" in result["warnings"]
    assert (extractor.cache / key / "page-2.png").exists()
    assert extractor.extract(data, ".pdf")[1]
    other = Extractor(tmp_path, {"ocr_language": "eng"})
    assert other.version != extractor.version


def test_word_and_slides_discovered_and_cached(env):
    service, _, _ = env
    doc = Document()
    doc.add_heading("Regression", 1)
    doc.add_paragraph("Ordinary least squares minimizes squared residuals.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "slope"
    table.cell(0, 1).text = "coefficient"
    doc.save(service.config.vault / "Study/documents/Lecture.docx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Regression lecture"
    slide.placeholders[1].text = "Residual variance"
    deck.save(service.config.vault / "Study/documents/Lecture.pptx")
    service.refresh()
    hits = service.search("Regression")
    assert {h["path"].split(".")[-1] for h in hits} == {"docx", "pptx"}
    assert any("paragraph" in h["locator"] for h in hits)
    assert any("slide" in h["locator"] for h in hits)
    assert service.refresh(full=True)["extracted"] == 0
    for row in service.db.execute("SELECT warnings FROM documents WHERE path LIKE '%.pptx'"):
        assert "visual_review_required" in json.loads(row[0])


def test_office_attachment_manifest(tmp_path):
    image = Image.new("RGB", (40, 40), "blue")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    stream.seek(0)
    doc = Document()
    doc.add_paragraph("Embedded diagram")
    doc.add_picture(stream)
    buffer = io.BytesIO()
    doc.save(buffer)
    extractor = Extractor(tmp_path)
    result, _, key = extractor.extract(buffer.getvalue(), ".docx")
    attachment = result["attachments"][0]
    assert attachment["package_part"].startswith("word/media/")
    assert (extractor.cache / key / attachment["artifact"]).exists()
