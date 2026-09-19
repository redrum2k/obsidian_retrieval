"""Content-addressed local extraction; no original file is passed to converters."""

import importlib.metadata
import io
import json
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import pymupdf
import yaml
from docx import Document
from pptx import Presentation

from .common import VaultError, canonical, digest

VERSION = "sections-v2"


def tool_version(command):
    if not shutil.which(command):
        return "unavailable"
    try:
        result = subprocess.run([command, "--version"], capture_output=True, timeout=10)
    except (OSError, subprocess.TimeoutExpired):
        return "unavailable"
    return digest(result.stdout + result.stderr)


class Extractor:
    def __init__(self, state, settings=None):
        self.cache = state / "extractions"
        self.cache.mkdir(parents=True, exist_ok=True)
        self.settings = settings or {}
        self.version = digest(
            canonical(
                {
                    "schema": VERSION,
                    "settings": self.settings,
                    "packages": {
                        p: importlib.metadata.version(p)
                        for p in ("PyMuPDF", "python-docx", "python-pptx", "PyYAML")
                    },
                    "tesseract": tool_version("tesseract"),
                    "soffice": tool_version("soffice"),
                }
            ).encode()
        )

    def extract(self, data, suffix):
        key = digest((digest(data) + suffix + self.version).encode())
        folder = self.cache / key
        result_path = folder / "result.json"
        if result_path.exists():
            return json.loads(result_path.read_text()), True, key
        folder.mkdir(exist_ok=True)
        sections, warnings, metadata = [], [], {}
        attachments = []
        if suffix in {".docx", ".pptx"}:
            prefix = "word/media/" if suffix == ".docx" else "ppt/media/"
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                for part in archive.namelist():
                    if part.startswith(prefix) and not part.endswith("/"):
                        blob = archive.read(part)
                        filename = "attachment-" + digest(blob)[:20] + Path(part).suffix
                        (folder / filename).write_bytes(blob)
                        attachments.append({"artifact": filename, "package_part": part})
        if suffix in {".md", ".markdown"}:
            text = data.decode("utf-8")
            lines = text.splitlines(keepends=True)
            if text.startswith("---\n"):
                end = text.find("\n---", 4)
                if end >= 0:
                    parsed = yaml.safe_load(text[4:end])
                    if isinstance(parsed, dict):
                        metadata = {
                            k: parsed[k] for k in ("title", "aliases", "author") if k in parsed
                        }
            heading, start, buf = [], 1, []
            for i, line in enumerate(lines, 1):
                match = re.match(r"^(#{1,6})\s+(.+)", line)
                if match and buf:
                    sections.append(
                        {
                            "heading": " / ".join(heading),
                            "text": "".join(buf),
                            "locator": {"line_start": start, "line_end": i - 1},
                        }
                    )
                    buf, start = [], i
                if match:
                    heading = heading[: len(match[1]) - 1] + [match[2].strip()]
                buf.append(line)
            if buf:
                sections.append(
                    {
                        "heading": " / ".join(heading),
                        "text": "".join(buf),
                        "locator": {"line_start": start, "line_end": len(lines)},
                    }
                )
        elif suffix == ".pdf":
            sections, warnings = self.pdf(data, folder)
        elif suffix == ".docx":
            doc = Document(io.BytesIO(data))
            heading = ""
            for i, para in enumerate(doc.paragraphs, 1):
                if para.style.name.startswith("Heading"):
                    heading = para.text
                if para.text:
                    sections.append(
                        {"heading": heading, "text": para.text, "locator": {"paragraph": i}}
                    )
            for i, table in enumerate(doc.tables, 1):
                sections.append(
                    {
                        "heading": "Table",
                        "text": "\n".join(
                            " | ".join(c.text for c in row.cells) for row in table.rows
                        ),
                        "locator": {"table": i},
                    }
                )
            warnings = ["visual_review_required"]
            self.office_render(data, suffix, folder, warnings)
        elif suffix == ".pptx":
            deck = Presentation(io.BytesIO(data))
            for i, slide in enumerate(deck.slides, 1):
                parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text)
                    if shape.has_table:
                        parts.extend(" | ".join(c.text for c in r.cells) for r in shape.table.rows)
                if slide.has_notes_slide:
                    parts.append(slide.notes_slide.notes_text_frame.text)
                sections.append(
                    {
                        "heading": slide.shapes.title.text if slide.shapes.title else "",
                        "text": "\n".join(parts),
                        "locator": {"slide": i},
                    }
                )
            warnings = ["visual_review_required"]
            self.office_render(data, suffix, folder, warnings)
            if (folder / "render.pdf").exists():
                rendered, pdf_warnings = self.pdf((folder / "render.pdf").read_bytes(), folder)
                warnings.extend(pdf_warnings)
                for i, section in enumerate(sections):
                    if not section["text"].strip() and i < len(rendered):
                        section["text"] = rendered[i]["text"]
        elif suffix in {".doc", ".ppt", ".odt", ".odp"}:
            self.office_render(data, suffix, folder, warnings)
            if not (folder / "render.pdf").exists():
                raise VaultError("unsupported_extraction", "Install LibreOffice for this format.")
            sections, warnings = self.pdf((folder / "render.pdf").read_bytes(), folder)
            warnings.append("visual_review_required")
        else:
            raise VaultError("unsupported_extraction", "No extractor for this format.")
        # Split long sections into continuation-addressable source character spans.
        chunks = []
        for section in sections:
            text = section["text"]
            for offset in range(0, max(1, len(text)), 1800):
                chunks.append(
                    {
                        **section,
                        "text": text[offset : offset + 1800],
                        "locator": {
                            **section["locator"],
                            "character_start": offset,
                            "character_end": min(offset + 1800, len(text)),
                        },
                    }
                )
        result = {
            "sections": chunks,
            "warnings": sorted(set(warnings)),
            "metadata": metadata,
            "attachments": attachments,
        }
        temp = folder / "result.tmp"
        temp.write_text(canonical(result))
        temp.replace(result_path)
        return result, False, key

    def pdf(self, data, folder):
        sections, warnings = [], []
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            for i, page in enumerate(doc, 1):
                text = page.get_text()
                # Cached renders provide evidence even when OCR is unavailable.
                image = folder / f"page-{i}.png"
                page.get_pixmap(dpi=150).save(image)
                if page.get_images() or page.get_drawings():
                    warnings.append("visual_review_required")
                if not text.strip():
                    warnings.append("visual_review_required")
                    if shutil.which("tesseract"):
                        result = subprocess.run(
                            [
                                "tesseract",
                                str(image),
                                "stdout",
                                "-l",
                                self.settings.get("ocr_language", "eng"),
                            ],
                            capture_output=True,
                            timeout=120,
                        )
                        if result.returncode:
                            raise VaultError(
                                "extraction_failed", "OCR failed; inspect local dependencies."
                            )
                        text = result.stdout.decode("utf-8")
                    else:
                        warnings.append("ocr_unavailable")
                sections.append({"heading": f"Page {i}", "text": text, "locator": {"page": i}})
        return sections, warnings

    def office_render(self, data, suffix, folder, warnings):
        if not shutil.which("soffice"):
            warnings.append("render_unavailable")
            return
        with tempfile.TemporaryDirectory(dir=folder) as scratch:
            root = Path(scratch)
            source = root / ("source" + suffix)
            source.write_bytes(data)
            result = subprocess.run(
                [
                    "soffice",
                    f"-env:UserInstallation={(root / 'profile').as_uri()}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(root),
                    str(source),
                ],
                capture_output=True,
                timeout=120,
            )
            pdf = root / "source.pdf"
            if result.returncode or not pdf.exists():
                raise VaultError("extraction_failed", "Office conversion failed.")
            shutil.copyfile(pdf, folder / "render.pdf")
            self.pdf(pdf.read_bytes(), folder)
